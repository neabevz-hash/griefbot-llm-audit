"""Build Word files from the Markdown masters with pandoc (shared by all papers).

Used by each paper's manuscript/build.py. The master Markdown may contain:
- <!-- author-block --> … <!-- /author-block -->        removed in the anonymized build;
- <!-- acknowledgements --> … <!-- /acknowledgements -->  removed in the anonymized build;
- placeholders such as ⟦OSF⟧, ⟦REPO⟧, ⟦DOI⟧               filled from a dict, or blinded.

Style: APA 7 professional manuscript — Times New Roman 12 pt, double spacing, 1-inch margins,
page numbers top right, headings by APA level (## centered bold, ### flush-left bold,
#### flush-left bold italic), tables single-spaced in 10 pt.
"""
from __future__ import annotations

import re
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PAGE_BREAK = '\n```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```\n'
BLIND_TEXT = "[link removed for anonymous review]"


def pandoc() -> str:
    exe = shutil.which("pandoc") or str(Path.home() / "scoop" / "shims" / "pandoc.exe")
    if not Path(exe).exists() and not shutil.which(exe):
        raise SystemExit("pandoc not found (scoop install pandoc)")
    return exe


def _font(style, size: float, bold: bool | None = None, italic: bool | None = None) -> None:
    f = style.font
    f.name, f.size, f.color.rgb = "Times New Roman", Pt(size), RGBColor(0, 0, 0)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(a), "Times New Roman")
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if fonts.get(qn(a)) is not None:
            del fonts.attrib[qn(a)]


def _spacing(style, line: float, before: float = 0, after: float = 0, indent: float | None = None,
             align=None, keep_next: bool = False) -> None:
    pf = style.paragraph_format
    pf.line_spacing, pf.space_before, pf.space_after = line, Pt(before), Pt(after)
    if indent is not None:
        pf.first_line_indent = Inches(indent)
    if align is not None:
        pf.alignment = align
    pf.keep_with_next = keep_next


def _page_number_header(section) -> None:
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.name, run.font.size = "Times New Roman", Pt(12)


def make_reference_docx(path: Path, double: bool = True) -> Path:
    """APA-style reference.docx for pandoc."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("wb") as fh:
        fh.write(subprocess.run([pandoc(), "--print-default-data-file", "reference.docx"],
                                check=True, capture_output=True).stdout)
    d = Document(str(path))
    line = 2.0 if double else 1.15
    st = d.styles
    for name in ("Normal", "Body Text", "First Paragraph", "Block Text", "Abstract", "Bibliography",
                 "Footnote Text", "Definition", "Definition Term", "Author", "Date", "Subtitle"):
        if name in [s.name for s in st]:
            _font(st[name], 12)
            _spacing(st[name], line)
    if double:   # APA manuscript: indented paragraphs, no space between them
        _spacing(st["Body Text"], line, indent=0.5)
        _spacing(st["First Paragraph"], line, indent=0.5)
    else:        # letters and supplements: block paragraphs
        _spacing(st["Body Text"], line, after=8, indent=0)
        _spacing(st["First Paragraph"], line, after=8, indent=0)
    _spacing(st["Bibliography"], line)
    st["Bibliography"].paragraph_format.left_indent = Inches(0.5)       # APA hanging indent
    st["Bibliography"].paragraph_format.first_line_indent = Inches(-0.5)
    _font(st["Compact"], 10)
    _spacing(st["Compact"], 1.0, indent=0)   # table cells: no first-line indent inherited from Body Text
    for name, (bold, italic, align) in {
        "Title": (True, False, WD_ALIGN_PARAGRAPH.CENTER),
        "Heading 1": (True, False, WD_ALIGN_PARAGRAPH.CENTER),
        "Heading 2": (True, False, WD_ALIGN_PARAGRAPH.CENTER),
        "Heading 3": (True, False, WD_ALIGN_PARAGRAPH.LEFT),
        "Heading 4": (True, True, WD_ALIGN_PARAGRAPH.LEFT),
        "Heading 5": (True, True, WD_ALIGN_PARAGRAPH.LEFT),
        "Abstract Title": (True, False, WD_ALIGN_PARAGRAPH.CENTER),
        "Table Caption": (False, False, WD_ALIGN_PARAGRAPH.LEFT),
        "Image Caption": (False, False, WD_ALIGN_PARAGRAPH.LEFT),
        "Caption": (False, False, WD_ALIGN_PARAGRAPH.LEFT),
    }.items():
        if name not in [s.name for s in st]:
            continue
        _font(st[name], 12, bold, italic)
        _spacing(st[name], line, align=align, keep_next=name.startswith("Heading"))
    # author block on title pages: centered, no indent
    _font(st["Author"], 12, bold=False, italic=False)
    _spacing(st["Author"], line, indent=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    # table notes: flush left, single-spaced, same size as the table text
    if "Table Note" not in [s.name for s in st]:
        tn = st.add_style("Table Note", WD_STYLE_TYPE.PARAGRAPH)
        tn.base_style = st["Normal"]
        _font(tn, 10)
        _spacing(tn, 1.0, before=3, after=12, indent=0)
    # tables: pandoc's "Table" style; keep the default grid-less look, single spacing via Compact
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(1))
    _page_number_header(sec)
    d.save(str(path))
    return path


def strip_block(md: str, tag: str) -> str:
    return re.sub(rf"<!--\s*{tag}\s*-->.*?<!--\s*/{tag}\s*-->\n?", "", md, flags=re.S)


def keep_block(md: str, tag: str) -> str:
    return re.sub(rf"<!--\s*/?{tag}\s*-->\n?", "", md)


def fill(md: str, values: dict[str, str | None], blinded: bool) -> str:
    def rep(m: re.Match) -> str:
        key = m.group(1)
        if blinded:
            return BLIND_TEXT
        v = values.get(key)
        return v if v else m.group(0)
    return re.sub(r"⟦([A-Za-z]+)⟧", rep, md)


def wrap_section(md: str, heading: str, style: str) -> str:
    """Give the paragraphs under a heading a custom Word style (e.g. Bibliography for hanging indents).
    The section ends at the next heading or table caption."""
    m = re.search(rf"(?m)^{re.escape(heading)}\s*$", md)
    if not m:
        return md
    n = re.search(r"(?m)^(#{1,6} |\*\*Table \d)", md[m.end():])
    end = m.end() + n.start() if n else len(md)
    body = md[m.end():end].strip("\n")
    return md[:m.end()] + f'\n\n::: {{custom-style="{style}"}}\n{body}\n:::\n\n' + md[end:]


def page_breaks(md: str, before: list[str]) -> str:
    """Insert a page break before every line that starts with one of the given prefixes."""
    out = []
    for line in md.splitlines():
        if any(line.startswith(b) for b in before):
            out.append(PAGE_BREAK)
        out.append(line)
    return "\n".join(out)


def break_before(docx_path: Path, patterns: list[str]) -> int:
    """Start a new page before every paragraph whose text matches one of the regexes
    (paragraph property, so no blank pages appear when a break falls at a page top)."""
    d = Document(str(docx_path))
    rx = [re.compile(p) for p in patterns]
    n = 0
    for par in d.paragraphs:
        if any(r.match(par.text.strip()) for r in rx):
            par.paragraph_format.page_break_before = True
            n += 1
    d.save(str(docx_path))
    return n


def wrap_lines(md: str, prefix: str, style: str) -> str:
    """Give every paragraph that starts with the prefix (e.g. '*Note.*') a custom Word style."""
    return re.sub(rf"(?m)^({re.escape(prefix)}.*)$",
                  lambda m: f'::: {{custom-style="{style}"}}\n{m.group(1)}\n:::', md)


def to_docx(md: str, out: Path, reference: Path, resource_dir: Path | None = None) -> Path:
    out.parent.mkdir(parents=True, exist_ok=True)
    src = out.with_suffix(".build.md")
    src.write_text(md, encoding="utf-8")
    cmd = [pandoc(), str(src), "-f", "markdown+pipe_tables+raw_attribute-smart", "-t", "docx",
           "--reference-doc", str(reference), "-o", str(out)]
    if resource_dir:
        cmd += ["--resource-path", str(resource_dir)]
    subprocess.run(cmd, check=True)
    src.unlink()
    return out


def word_count(md: str) -> int:
    md = re.sub(r"<!--.*?-->", " ", md, flags=re.S)
    md = re.sub(r"```\{=openxml\}.*?```", " ", md, flags=re.S)
    return len(re.findall(r"\S*[A-Za-z0-9]\S*", md.replace("|", " ").replace("**", " ")))
